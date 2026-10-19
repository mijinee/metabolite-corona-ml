"""
External validation of HILIC class predictions against LIPID MAPS.

Input:  LipidMaps_Feature_Mapping.xlsx
Output: LipidMaps_Accuracy.xlsx
"""

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from config import HILIC_TRACK1, HILIC_TRACK2, LIPIDMAPS_MAPPING, OUT_ML_EVAL

OUTPUT_DIR = OUT_ML_EVAL / "lipidmaps_validation"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
EXCEL_OUTPUT = OUTPUT_DIR / "LipidMaps_Accuracy.xlsx"

ABBREV_TO_CLASS = {
    "FA": "Fatty Acids & Amides", "NAT": "Fatty Acids & Amides", "NAE": "Fatty Acids & Amides",
    "FAHFA": "Fatty Acids & Amides", "FOH": "Fatty Acids & Amides", "WE": "Fatty Acids & Amides",
    "CAR": "Acylcarnitines",
    "PA": "Phospholipids", "PE": "Phospholipids", "PC": "Phospholipids",
    "PS": "Phospholipids", "PG": "Phospholipids", "PI": "Phospholipids",
    "LPA": "Lysophospholipids", "LPE": "Lysophospholipids", "LPC": "Lysophospholipids",
    "LPS": "Lysophospholipids", "LPG": "Lysophospholipids", "LPI": "Lysophospholipids",
    "ST": "Complex Lipids", "SFE": "Complex Lipids",
    "SM": "Complex Lipids", "Cer": "Complex Lipids", "CerP": "Complex Lipids",
    "CerPE": "Complex Lipids", "HexCer": "Complex Lipids", "Hex2Cer": "Complex Lipids",
    "SPB": "Complex Lipids", "SPBP": "Complex Lipids",
}

LIPID_MAPS_VALIDATABLE = {"Phospholipids", "Lysophospholipids", "Fatty Acids & Amides",
                          "Complex Lipids", "Acylcarnitines"}

TRACK_MAP = {**{c: "Track 1" for c in HILIC_TRACK1}, **{c: "Track 2" for c in HILIC_TRACK2}}


def get_abbrev(name):
    m = re.match(r"^([A-Za-z0-9\-]+)\s", str(name).strip())
    return m.group(1) if m else str(name).strip()


# Load hits
raw = pd.read_excel(LIPIDMAPS_MAPPING, sheet_name="LipidMaps_Mapping")
matched = raw[raw["LipidMaps_Name"].notna()].copy()
matched["Abbrev"] = matched["LipidMaps_Name"].apply(get_abbrev)
matched["LipidMaps_Class"] = matched["Abbrev"].map(ABBREV_TO_CLASS)

unmapped = matched[matched["LipidMaps_Class"].isna()]["Abbrev"].unique()
if len(unmapped):
    print("WARNING - unmapped abbreviations found:", list(unmapped))
print("Matched hit rows:", len(matched), "| Unique matched features:", matched["FeatureID"].nunique())

# Feature-level top-hit match
rows = []
for fid, grp in matched.groupby("FeatureID"):
    predicted  = grp["Predicted_RandomForest"].iloc[0]
    min_delta  = grp["Mass_Difference"].round(4).min()
    top_hits   = grp[grp["Mass_Difference"].round(4) == min_delta]
    top_classes = set(top_hits["LipidMaps_Class"].dropna())
    all_classes = set(grp["LipidMaps_Class"].dropna())
    rows.append({
        "FeatureID": fid,
        "Predicted_RandomForest": predicted,
        "Confidence_RandomForest": grp["Confidence_RandomForest"].iloc[0],
        "Track": TRACK_MAP.get(predicted, "Unassigned"),
        "LipidMaps_Validatable_Class": predicted in LIPID_MAPS_VALIDATABLE,
        "N_Hits": len(grp),
        "N_Tied_Top_Hits": len(top_hits),
        "Min_Mass_Delta": min_delta,
        "LipidMaps_TopHit_Classes": ", ".join(sorted(top_classes)),
        "LipidMaps_TopHit_Match": predicted in top_classes,
        "LipidMaps_AnyHit_Classes": ", ".join(sorted(all_classes)),
        "LipidMaps_AnyHit_Match": predicted in all_classes,
    })
feature_level = pd.DataFrame(rows).sort_values("FeatureID")

scoreable = feature_level[(feature_level["Predicted_RandomForest"] != "Low Confidence") &
                          (feature_level["LipidMaps_Validatable_Class"])].copy()

# Track summary
track_rows = []
for track in ["Track 1", "Track 2"]:
    sub = scoreable[scoreable["Track"] == track]
    n, n_match = len(sub), int(sub["LipidMaps_TopHit_Match"].sum())
    track_rows.append({"Track": track, "N_Scoreable_Features": n, "N_TopHit_Match": n_match,
                       "LipidMaps_TopHit_Accuracy_%": round(n_match / n * 100, 1) if n else np.nan})
n, n_match = len(scoreable), int(scoreable["LipidMaps_TopHit_Match"].sum())
track_rows.append({"Track": "Overall", "N_Scoreable_Features": n, "N_TopHit_Match": n_match,
                   "LipidMaps_TopHit_Accuracy_%": round(n_match / n * 100, 1) if n else np.nan})
track_summary = pd.DataFrame(track_rows)

# Class breakdown
class_rows = []
confident = feature_level[feature_level["Predicted_RandomForest"] != "Low Confidence"]
for cls, grp in confident.groupby("Predicted_RandomForest"):
    n, n_match = len(grp), int(grp["LipidMaps_TopHit_Match"].sum())
    class_rows.append({"Predicted_RandomForest": cls, "Track": TRACK_MAP.get(cls, "Unassigned"),
                       "LipidMaps_Validatable": cls in LIPID_MAPS_VALIDATABLE,
                       "N": n, "N_TopHit_Match": n_match,
                       "LipidMaps_TopHit_Accuracy_%": round(n_match / n * 100, 1)
                       if cls in LIPID_MAPS_VALIDATABLE else np.nan})
class_summary = pd.DataFrame(class_rows).sort_values(["Track", "N"], ascending=[True, False])

# Confusion matrix (predicted class vs LIPID MAPS top-hit class)
exploded = confident.assign(TopHitClass=confident["LipidMaps_TopHit_Classes"].str.split(", ")) \
                    .explode("TopHitClass").reset_index(drop=True)
confusion = pd.crosstab([exploded["Track"], exploded["Predicted_RandomForest"]], exploded["TopHitClass"])
row_order = class_summary[["Track", "Predicted_RandomForest"]].apply(tuple, axis=1).tolist()
confusion = confusion.reindex(row_order)

# Confidence calibration
bins = np.arange(0.3, 1.1, 0.1)
scoreable["Confidence_Bin"] = pd.cut(scoreable["Confidence_RandomForest"], bins=bins, right=False)
calibration = scoreable.groupby("Confidence_Bin", observed=False).agg(
    N=("LipidMaps_TopHit_Match", "size"),
    N_Match=("LipidMaps_TopHit_Match", "sum")).reset_index()
calibration = calibration[calibration["N"] > 0].copy()
calibration["Confidence_Bin"] = calibration["Confidence_Bin"].astype(str)
calibration["Empirical_Accuracy_%"] = (calibration["N_Match"] / calibration["N"] * 100).round(1)

# Export
readme_lines = [
    ("LipidMaps Top-Hit Accuracy - Methodology", None),
    ("", None),
    ("Primary metric:", "Among RF-predicted features with >=1 LipidMaps mass hit, the % where the "
                        "predicted class matches the class of the closest-mass hit(s)."),
    ("", "Ties at the minimum mass difference are all counted as top hits."),
    ("Exclusions:", "'Low Confidence' predictions, and non-lipid predicted classes (not represented "
                    "in LIPID MAPS; shown in the confusion matrix only)."),
    ("Coverage:", f"{matched['FeatureID'].nunique()} Predict-set features had >=1 hit "
                  f"within +/-0.05 Da and matched adduct polarity."),
    ("Predict set:", "Split == 'Predict' in the class prediction output."),
    ("Track 1:", ", ".join(HILIC_TRACK1)),
    ("Track 2:", ", ".join(HILIC_TRACK2)),
]

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
BODY_FONT   = Font(name="Arial")


def write_df(ws, df):
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    for i, (_, r) in enumerate(df.iterrows(), start=2):
        for j, col in enumerate(df.columns, start=1):
            ws.cell(row=i, column=j, value=r[col]).font = BODY_FONT
    for j, col in enumerate(df.columns, start=1):
        longest = int(df[col].astype(str).str.len().max()) if len(df) else 10
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(40, longest + 2, len(str(col)) + 4))
    ws.freeze_panes = ws.cell(row=2, column=1)


wb = Workbook()
ws = wb.active
ws.title = "README"
for i, (label, text) in enumerate(readme_lines, start=1):
    if label:
        ws.cell(row=i, column=1, value=label).font = Font(name="Arial", bold=True)
    if text:
        ws.cell(row=i, column=2, value=text).font = BODY_FONT
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 100

cm_reset = confusion.reset_index()
cm_reset.columns = ["Track", "Predicted_RandomForest"] + list(confusion.columns)
for name, df_out in [("Track_Summary", track_summary), ("Class_Breakdown", class_summary),
                     ("Confusion_Matrix", cm_reset), ("Confidence_Calibration", calibration),
                     ("Feature_Level", feature_level)]:
    write_df(wb.create_sheet(name), df_out)
wb.save(EXCEL_OUTPUT)

print(track_summary.to_string(index=False))
print(f"Saved: {EXCEL_OUTPUT}")
